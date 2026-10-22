# app/api/v1/students/import_helper.py

from fastapi import HTTPException, UploadFile
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from io import BytesIO
import logging
from datetime import datetime

logger = logging.getLogger(__name__)


# ============================================================
# ✅ GENERATE REGISTRATION NUMBER (AUTO) - KWA DARASA PEKEE
# ============================================================
def generate_registration_number(db, class_id: int, level: str = "secondary") -> str:
    """
    ✅ Zalisha Namba ya Usajili ya kipekee KWA DARASA PEKEE.
    """
    from app.models.student import Student
    
    prefix = "P" if level.lower() == "primary" else "S"
    
    last_student = db.query(Student).filter(
        Student.class_id == class_id,
        Student.roll_number.like(f"{prefix}.%")
    ).order_by(Student.id.desc()).first()
    
    if last_student and last_student.roll_number:
        try:
            last_num = int(last_student.roll_number.split('.')[1])
            new_num = last_num + 1
        except:
            new_num = 1
    else:
        new_num = 1
    
    return f"{prefix}.{new_num:04d}"


# ============================================================
# ✅ PARSE PReM EXCEL (SERIKALI)
# ============================================================
def parse_prem_excel(file_content: bytes) -> list:
    """
    ✅ Chambua Excel ya PReM (serikali) na urudishe list ya wanafunzi.
    
    Inatambua:
    - Jina la Kwanza, Kati, Mwisho → Jina Kamili
    - Jinsia (1=M, 2=F)
    - Darasa
    - Mkondo (kama upo)
    - Jina la Mzazi (Kwanza + Kati + Mwisho)
    - Simu ya Mzazi
    """
    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True)
        sheet = workbook.active
    except Exception as e:
        raise HTTPException(400, f"Imeshindwa kusoma Excel: {str(e)}")
    
    # ✅ Pata headers kwenye rows 1-15 (PReM ina headers ngumu)
    header_row = None
    headers = []
    
    for row_idx in range(1, 16):
        row_values = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[row_idx]]
        
        has_jina = any("jina" in v for v in row_values)
        has_jinsia = any("jinsia" in v or "jinsi" in v for v in row_values)
        
        if has_jina and has_jinsia:
            header_row = row_idx
            headers = row_values
            break
    
    if header_row is None:
        return None  # Sio PReM Excel
    
    logger.info(f"📋 PReM header row: {header_row}")
    logger.info(f"📋 PReM Headers: {headers}")
    
    # ✅ Tafuta index ya kila column
    col_indices = {}
    for i, header in enumerate(headers):
        # Jina la Mwanafunzi (Kwanza, Kati, Mwisho)
        if 'jina' in header and 'kwanza' in header and 'mzazi' not in header:
            col_indices['first_name'] = i
        elif 'jina' in header and 'kati' in header and 'mzazi' not in header:
            col_indices['middle_name'] = i
        elif 'jina' in header and 'mwisho' in header and 'mzazi' not in header:
            col_indices['last_name'] = i
        
        # Jinsia
        elif 'jinsia' in header or 'jinsi' in header:
            col_indices['sex'] = i
        
        # Darasa
        elif 'darasa' in header:
            col_indices['class'] = i
        
        # Mkondo
        elif 'mkondo' in header or 'stream' in header:
            col_indices['stream'] = i
        
        # ✅ Jina la Mzazi (Kwanza, Kati, Mwisho)
        elif 'jina' in header and 'kwanza' in header and ('mzazi' in header or 'baba' in header or 'mlezi' in header):
            col_indices['parent_first_name'] = i
        elif 'jina' in header and 'kati' in header and ('mzazi' in header or 'baba' in header or 'mlezi' in header):
            col_indices['parent_middle_name'] = i
        elif 'jina' in header and 'mwisho' in header and ('mzazi' in header or 'baba' in header or 'mlezi' in header):
            col_indices['parent_last_name'] = i
        
        # ✅ Simu ya Mzazi
        elif 'simu' in header:
            col_indices['parent_phone'] = i
    
    logger.info(f"📍 PReM Column indices: {col_indices}")
    
    if 'first_name' not in col_indices:
        return None
    
    # ✅ Chambua rows (kuanzia header_row + 1)
    students = []
    errors = []
    
    for row_num, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not row or not any(row):
            continue
        
        try:
            student = {}
            
            # ✅ Jina Kamili = Kwanza + Kati + Mwisho
            first = str(row[col_indices['first_name']]).strip() if 'first_name' in col_indices and row[col_indices['first_name']] else ""
            middle = str(row[col_indices['middle_name']]).strip() if 'middle_name' in col_indices and row[col_indices['middle_name']] else ""
            last = str(row[col_indices['last_name']]).strip() if 'last_name' in col_indices and row[col_indices['last_name']] else ""
            
            full_name = " ".join(filter(None, [first, middle, last])).strip()
            
            if not full_name:
                errors.append(f"Row {row_num}: Jina halipo")
                continue
            
            student['name'] = full_name
            
            # ✅ Jinsia (1=M, 2=F, au M/F)
            sex_idx = col_indices.get('sex')
            if sex_idx is not None and row[sex_idx]:
                sex = str(row[sex_idx]).strip().upper()
                if sex in ['1', 'M', 'ME', 'MALE', 'MUME', 'MVULANA']:
                    student['sex'] = 'M'
                elif sex in ['2', 'F', 'KE', 'FEMALE', 'MKE', 'MSICHANA']:
                    student['sex'] = 'F'
                else:
                    student['sex'] = sex[:1]
            else:
                errors.append(f"Row {row_num}: Jinsia haipo")
                continue
            
            # ✅ Darasa
            class_idx = col_indices.get('class')
            if class_idx is not None and row[class_idx]:
                student['class'] = str(row[class_idx]).strip()
            else:
                student['class'] = "Darasa la 1"  # Default
            
            # ✅ Mkondo (kama upo)
            stream_idx = col_indices.get('stream')
            if stream_idx is not None and row[stream_idx]:
                student['stream'] = str(row[stream_idx]).strip()
            
            # ✅ Jina la Mzazi (Kwanza + Kati + Mwisho)
            p_first = str(row[col_indices['parent_first_name']]).strip() if 'parent_first_name' in col_indices and row[col_indices['parent_first_name']] else ""
            p_middle = str(row[col_indices['parent_middle_name']]).strip() if 'parent_middle_name' in col_indices and row[col_indices['parent_middle_name']] else ""
            p_last = str(row[col_indices['parent_last_name']]).strip() if 'parent_last_name' in col_indices and row[col_indices['parent_last_name']] else ""
            
            parent_full_name = " ".join(filter(None, [p_first, p_middle, p_last])).strip()
            
            if parent_full_name:
                student['parent_name'] = parent_full_name
            
            # ✅ Simu ya Mzazi
            phone_idx = col_indices.get('parent_phone')
            if phone_idx is not None and row[phone_idx]:
                student['parent_phone'] = str(row[phone_idx]).strip()
            
            students.append(student)
            
        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")
    
    logger.info(f"✅ PReM Parsed {len(students)} students, {len(errors)} errors")
    
    if not students:
        raise HTTPException(400, f"Hakuna wanafunzi sahihi. Errors: {errors[:5]}")
    
    return students


# ============================================================
# ✅ PARSE EXCEL (TEMPLATE YETU AU PReM)
# ============================================================
def parse_excel_students(file_content: bytes) -> list:
    """
    ✅ Chambua Excel — inajaribu PReM kwanza, kisha template yetu.
    """
    # ✅ Jaribu PReM kwanza
    prem_students = parse_prem_excel(file_content)
    if prem_students is not None:
        logger.info(f"✅ PReM Excel detected! Using PReM parser.")
        return prem_students
    
    # ✅ Kama sio PReM, tumia parser ya kawaida
    logger.info(f"📋 Using standard parser.")
    
    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True)
        sheet = workbook.active
    except Exception as e:
        raise HTTPException(400, f"Imeshindwa kusoma Excel: {str(e)}")
    
    # Pata headers
    headers = []
    for cell in sheet[1]:
        if cell.value:
            headers.append(str(cell.value).strip().lower())
    
    logger.info(f"📋 Standard Headers: {headers}")
    
    # Kagua headers muhimu
    required_keywords = ['jina', 'jinsia', 'darasa']
    missing = []
    for keyword in required_keywords:
        if not any(keyword in h for h in headers):
            missing.append(keyword)
    
    if missing:
        raise HTTPException(
            400,
            f"Excel haina columns: {', '.join(missing)}. "
            f"Iliyopo: {', '.join(headers)}"
        )
    
    # Tafuta index ya kila column
    col_indices = {}
    for i, header in enumerate(headers):
        if 'jina' in header and ('mzazi' in header or 'baba' in header or 'mlezi' in header):
            col_indices['parent_name'] = i
        elif 'simu' in header and ('mzazi' in header or 'baba' in header or 'mlezi' in header):
            col_indices['parent_phone'] = i
        elif 'jina' in header:
            col_indices['name'] = i
        elif 'jinsia' in header or 'jinsi' in header or 'sex' in header:
            col_indices['sex'] = i
        elif 'darasa' in header or 'class' in header:
            col_indices['class'] = i
        elif 'mkondo' in header or 'stream' in header:
            col_indices['stream'] = i
        elif 'usajili' in header or 'roll' in header or 'namba' in header:
            col_indices['roll_number'] = i
    
    logger.info(f"📍 Standard Column indices: {col_indices}")
    
    # Chambua rows
    students = []
    errors = []
    
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue
        
        try:
            student = {}
            
            # Jina (lazima)
            name_idx = col_indices.get('name')
            if name_idx is not None and row[name_idx]:
                student['name'] = str(row[name_idx]).strip()
            else:
                errors.append(f"Row {row_num}: Jina halipo")
                continue
            
            # Jinsia (lazima)
            sex_idx = col_indices.get('sex')
            if sex_idx is not None and row[sex_idx]:
                sex = str(row[sex_idx]).strip().upper()
                if sex in ['1', 'M', 'ME', 'MALE', 'MUME', 'MVULANA']:
                    student['sex'] = 'M'
                elif sex in ['2', 'F', 'KE', 'FEMALE', 'MKE', 'MSICHANA']:
                    student['sex'] = 'F'
                else:
                    student['sex'] = sex[:1]
            else:
                errors.append(f"Row {row_num}: Jinsia haipo")
                continue
            
            # Darasa (lazima)
            class_idx = col_indices.get('class')
            if class_idx is not None and row[class_idx]:
                student['class'] = str(row[class_idx]).strip()
            else:
                errors.append(f"Row {row_num}: Darasa halipo")
                continue
            
            # Optional fields
            if 'roll_number' in col_indices and row[col_indices['roll_number']]:
                student['roll_number'] = str(row[col_indices['roll_number']]).strip()
            
            if 'stream' in col_indices and row[col_indices['stream']]:
                student['stream'] = str(row[col_indices['stream']]).strip()
            
            if 'parent_name' in col_indices and row[col_indices['parent_name']]:
                student['parent_name'] = str(row[col_indices['parent_name']]).strip()
            
            if 'parent_phone' in col_indices and row[col_indices['parent_phone']]:
                student['parent_phone'] = str(row[col_indices['parent_phone']]).strip()
            
            students.append(student)
            
        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")
    
    logger.info(f"✅ Standard Parsed {len(students)} students, {len(errors)} errors")
    
    if not students:
        raise HTTPException(
            400,
            f"Hakuna wanafunzi sahihi. Errors: {errors[:5]}"
        )
    
    return students


# ============================================================
# ✅ IMPORT TO DATABASE
# ============================================================
def import_students_to_db(
    students_data: list,
    school_id: int,
    db: Session,
    SchoolClass,
    Stream,
    Student,
    level: str = "secondary",
) -> dict:
    """
    ✅ Ingiza wanafunzi kwenye database
    """
    imported = 0
    failed = 0
    errors = []
    
    for data in students_data:
        try:
            # ----------------------------------------
            # 1. Tafuta au unda darasa
            # ----------------------------------------
            class_obj = db.query(SchoolClass).filter(
                SchoolClass.school_id == school_id,
                SchoolClass.name.ilike(f"%{data['class']}%")
            ).first()
            
            if not class_obj:
                class_obj = SchoolClass(
                    school_id=school_id,
                    name=data['class'],
                )
                db.add(class_obj)
                db.flush()
                logger.info(f"➕ Created class: {data['class']}")
            
            # ----------------------------------------
            # 2. Tafuta au unda mkondo
            # ----------------------------------------
            stream_obj = None
            if data.get('stream'):
                stream_obj = db.query(Stream).filter(
                    Stream.class_id == class_obj.id,
                    Stream.name.ilike(f"%{data['stream']}%")
                ).first()
                
                if not stream_obj:
                    stream_obj = Stream(
                        class_id=class_obj.id,
                        name=data['stream'],
                        school_id=school_id,
                    )
                    db.add(stream_obj)
                    db.flush()
                    logger.info(f"➕ Created stream: {data['stream']}")
            
            # ----------------------------------------
            # 3. Zalisha namba ya usajili AUTOMATICALLY
            # ----------------------------------------
            roll_number = generate_registration_number(
                db=db,
                class_id=class_obj.id,
                level=level,
            )
            
            # ----------------------------------------
            # 4. Kagua kama mwanafunzi yupo tayari
            # ----------------------------------------
            existing = db.query(Student).filter(
                Student.school_id == school_id,
                Student.name.ilike(data['name']),
                Student.class_id == class_obj.id,
            ).first()
            
            if existing:
                logger.info(f"⏭️ Skipped (already exists): {data['name']}")
                failed += 1
                errors.append(f"{data['name']}: Tayari yupo kwenye darasa hili")
                continue
            
            # ----------------------------------------
            # 5. Unda mwanafunzi
            # ----------------------------------------
            new_student = Student(
                name=data['name'],
                sex=data['sex'],
                school_id=school_id,
                class_id=class_obj.id,
                stream_id=stream_obj.id if stream_obj else None,
                roll_number=roll_number,
                father_name=data.get('parent_name'),
                father_phone=data.get('parent_phone'),
                enrollment_date=datetime.utcnow(),
            )
            db.add(new_student)
            db.flush()
            imported += 1
            
            logger.info(f"✅ Imported: {data['name']} → {roll_number}")
            
        except Exception as e:
            failed += 1
            errors.append(f"{data.get('name', 'N/A')}: {str(e)}")
            logger.error(f"❌ Failed to import {data.get('name')}: {e}")
    
    try:
        db.commit()
        logger.info(f"✅ Committed {imported} students")
    except Exception as e:
        db.rollback()
        logger.error(f"❌ Commit failed: {e}")
        raise HTTPException(500, f"Database commit failed: {str(e)}")
    
    return {
        "imported": imported,
        "failed": failed,
        "errors": errors[:10],
        "total": len(students_data),
    }


# ============================================================
# ✅ CREATE TEMPLATE (TUPU - BILA SAMPLE)
# ============================================================
def create_import_template(level: str = "secondary") -> BytesIO:
    """
    ✅ Unda Excel template ya ku-import wanafunzi.
    TEMPLATE TUPU — Wakuu watajaza majina yao halisi.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Wanafunzi"
    
    headers = [
        'Jina Kamili',
        'Jinsia (M/F)',
        'Darasa',
        'Mkondo',
        'Jina la Mzazi',
        'Simu ya Mzazi'
    ]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # ✅ HAKUNA SAMPLE — Template tupu
    # Wakuu watajaza majina yao halisi.
    
    widths = [25, 15, 20, 15, 25, 20]
    for i, width in enumerate(widths):
        col_letter = chr(65 + i)
        ws.column_dimensions[col_letter].width = width
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer